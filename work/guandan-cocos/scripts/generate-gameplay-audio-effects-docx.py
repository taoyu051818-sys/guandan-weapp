#!/usr/bin/env python3
"""Generate the game-flow audio/effects review document with captured screenshots."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "陵水掼蛋牌局音效与动效说明.docx"
SHOT_DIR = ROOT / "docs" / "assets" / "game-flow"

NAVY = "123B46"
TEAL = "1C6B68"
GOLD = "D5A738"
PALE = "EAF4F2"
PALE_GOLD = "FBF3DD"
PALE_RED = "FBECE8"
INK = RGBColor(30, 48, 54)
MUTED = RGBColor(91, 108, 112)
WHITE = RGBColor(255, 255, 255)
FONT = "STSong"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_width(cell, width_cm: float) -> None:
    cell.width = Cm(width_cm)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_cm * 567)))
    tc_w.set(qn("w:type"), "dxa")


def set_table_grid_widths(table, widths_cm: list[float]) -> None:
    for idx, width_cm in enumerate(widths_cm):
        if idx < len(table.columns):
            table.columns[idx].width = Cm(width_cm)
    grid_cols = list(table._tbl.tblGrid.gridCol_lst)
    for idx, width_cm in enumerate(widths_cm):
        if idx < len(grid_cols):
            grid_cols[idx].set(qn("w:w"), str(int(width_cm * 567)))


def set_cell_left_border(cell, color: str, size: int = 26) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    left = borders.find(qn("w:left"))
    if left is None:
        left = OxmlElement("w:left")
        borders.append(left)
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), str(size))
    left.set(qn("w:space"), "0")
    left.set(qn("w:color"), color)


def set_run_font(run, size: float | None = None, bold: bool | None = None, color: RGBColor | None = None) -> None:
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def style_paragraph(paragraph, size=10.2, color=INK, bold=False, line=1.18, after=4) -> None:
    paragraph.paragraph_format.line_spacing = line
    paragraph.paragraph_format.space_after = Pt(after)
    for run in paragraph.runs:
        set_run_font(run, size=size, bold=bold, color=color)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_run_font(run, size=8.5, color=MUTED)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])
    tail = paragraph.add_run(" 页")
    set_run_font(tail, size=8.5, color=MUTED)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(10 if level == 1 else 7)
    p.paragraph_format.space_after = Pt(5)
    if level == 1:
        p.paragraph_format.left_indent = Cm(0)
        r = p.add_run(text)
        set_run_font(r, size=17, bold=True, color=RGBColor(18, 59, 70))
        border = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "12")
        bottom.set(qn("w:space"), "5")
        bottom.set(qn("w:color"), GOLD)
        border.append(bottom)
        p._p.get_or_add_pPr().append(border)
    else:
        r = p.add_run(text)
        set_run_font(r, size=12.5, bold=True, color=RGBColor(28, 107, 104))


def add_bullet(doc: Document, text: str, color=INK) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Cm(0.6)
    p.paragraph_format.first_line_indent = Cm(-0.2)
    p.add_run(text)
    style_paragraph(p, size=9.7, color=color, after=2)


def add_callout(doc: Document, title: str, body: str, fill=PALE_GOLD, accent=GOLD) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_grid_widths(table, [16.9])
    cell = table.cell(0, 0)
    set_cell_width(cell, 16.9)
    set_cell_shading(cell, fill)
    set_cell_left_border(cell, accent)
    set_cell_margins(cell, top=130, start=190, bottom=130, end=150)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title)
    set_run_font(r, size=10.5, bold=True, color=RGBColor(18, 59, 70))
    p2 = cell.add_paragraph(body)
    style_paragraph(p2, size=9.3, color=INK, line=1.15, after=0)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float], font_size=8.2) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    set_table_grid_widths(table, widths)
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_width(cell, widths[idx])
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell, top=85, start=90, bottom=85, end=90)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(header)
        set_run_font(r, size=8.7, bold=True, color=WHITE)
    set_repeat_table_header(table.rows[0])
    for ridx, values in enumerate(rows):
        row = table.add_row()
        for idx, value in enumerate(values):
            cell = row.cells[idx]
            set_cell_width(cell, widths[idx])
            set_cell_margins(cell, top=75, start=85, bottom=75, end=85)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if ridx % 2 == 1:
                set_cell_shading(cell, "F5F8F7")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.08
            r = p.add_run(value)
            set_run_font(r, size=font_size, color=INK, bold=(idx == 0))
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_flow_strip(doc: Document, labels: list[str]) -> None:
    cols = len(labels) * 2 - 1
    table = doc.add_table(rows=1, cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    grid_widths = [2.3 if idx % 2 == 0 else 0.45 for idx in range(cols)]
    set_table_grid_widths(table, grid_widths)
    for idx, label in enumerate(labels):
        cell = table.cell(0, idx * 2)
        set_cell_width(cell, 2.3)
        set_cell_shading(cell, PALE)
        set_cell_margins(cell, top=120, start=70, bottom=120, end=70)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(label)
        set_run_font(r, size=8.5, bold=True, color=RGBColor(28, 107, 104))
        if idx < len(labels) - 1:
            arrow = table.cell(0, idx * 2 + 1)
            set_cell_width(arrow, 0.45)
            p2 = arrow.paragraphs[0]
            p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r2 = p2.add_run("→")
            set_run_font(r2, size=12, bold=True, color=RGBColor(213, 167, 56))
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_picture(doc: Document, filename: str, caption: str, width_inches=6.55) -> None:
    path = SHOT_DIR / filename
    if not path.exists():
        raise FileNotFoundError(path)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(path), width=Inches(width_inches))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(6)
    r = cap.add_run(caption)
    set_run_font(r, size=8.5, color=MUTED)


def add_picture_pair(doc: Document, items: list[tuple[str, str]]) -> None:
    table = doc.add_table(rows=2, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_grid_widths(table, [8.25, 8.25])
    for idx, (filename, caption) in enumerate(items):
        col = idx % 2
        pic_cell = table.cell(0, col)
        cap_cell = table.cell(1, col)
        set_cell_width(pic_cell, 8.25)
        set_cell_width(cap_cell, 8.25)
        set_cell_margins(pic_cell, top=25, start=35, bottom=25, end=35)
        set_cell_margins(cap_cell, top=25, start=60, bottom=90, end=60)
        p = pic_cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(SHOT_DIR / filename), width=Inches(3.12))
        cp = cap_cell.paragraphs[0]
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cr = cp.add_run(caption)
        set_run_font(cr, size=8.1, color=MUTED)


def setup_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.65)
    section.bottom_margin = Cm(1.55)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    section.header_distance = Cm(0.7)
    section.footer_distance = Cm(0.7)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(10.2)
    normal.font.color.rgb = INK
    normal.paragraph_format.line_spacing = 1.18
    normal.paragraph_format.space_after = Pt(4)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = header.add_run("陵水掼蛋  ·  牌局体验规格")
    set_run_font(run, size=8.5, bold=True, color=RGBColor(28, 107, 104))
    add_page_number(section.footer.paragraphs[0])
    return doc


def build() -> None:
    doc = setup_document()

    # Cover
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(48)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("陵水掼蛋")
    set_run_font(r, size=19, bold=True, color=RGBColor(28, 107, 104))
    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(10)
    r = title.add_run("牌局流程、音效与动效说明")
    set_run_font(r, size=30, bold=True, color=RGBColor(18, 59, 70))
    sub = doc.add_paragraph()
    sub.paragraph_format.space_after = Pt(25)
    r = sub.add_run("当前完成度 · 截图材料 · 缺陷与人工验收基线")
    set_run_font(r, size=12.5, color=RGBColor(213, 167, 56))

    add_callout(
        doc,
        "本轮结论",
        "当前能稳定辨识的牌面报读主要为部分单张和对子。其他牌型即使已有源码配置，也要等实际听感验收后才能标为完成。联机增量动作目前还存在被误判为恢复快照的问题，会跳过音效、飞牌和主特效。",
        fill=PALE_GOLD,
        accent=GOLD,
    )
    meta = doc.add_table(rows=4, cols=2)
    meta.alignment = WD_TABLE_ALIGNMENT.LEFT
    meta.autofit = False
    set_table_grid_widths(meta, [3.5, 12.7])
    pairs = [
        ("文档版本", "2026-08-04"),
        ("对应构建", "Web 测试版 · http://127.0.0.1:7488/"),
        ("截图规格", "1280 × 720"),
        ("设计预设", "Standard Business Brief · 研发评审稿"),
    ]
    for idx, (key, value) in enumerate(pairs):
        set_cell_width(meta.cell(idx, 0), 3.5)
        set_cell_width(meta.cell(idx, 1), 12.7)
        set_cell_shading(meta.cell(idx, 0), PALE)
        for cell in meta.rows[idx].cells:
            set_cell_margins(cell, top=105, start=120, bottom=105, end=120)
        p1 = meta.cell(idx, 0).paragraphs[0]
        r1 = p1.add_run(key)
        set_run_font(r1, size=9.2, bold=True, color=RGBColor(28, 107, 104))
        p2 = meta.cell(idx, 1).paragraphs[0]
        r2 = p2.add_run(value)
        set_run_font(r2, size=9.2, color=INK)
    doc.add_paragraph()
    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(70)
    r = note.add_run("内部研发与测试文档 · 截图取自当前本地构建")
    set_run_font(r, size=8.8, color=MUTED)
    doc.add_page_break()

    # Coverage
    add_heading(doc, "1  验收口径与牌型覆盖", 1)
    p = doc.add_paragraph("本文把“已实测”“源码已配置，待复验”“待实现”严格分开。资源存在或代码存在不能替代真机听感验收。")
    style_paragraph(p, size=10, color=INK)
    add_bullet(doc, "通用出牌声：短促落牌声，不说明牌型。")
    add_bullet(doc, "牌面报读：例如“对五”“顺子”等可辨识语音。")
    add_bullet(doc, "牌型事件音效：例如炸弹爆破声，不一定包含人声报型。")
    coverage = [
        ["单张", "通用声；部分点数有报读", "260ms 飞牌、落地光环", "部分实测"],
        ["对子", "通用声；部分点数有报读", "280ms 飞牌、标签、流光", "部分实测"],
        ["三张", "无专用报型", "飞牌、标签、流光", "音频待补"],
        ["顺子", "授权语音已接入", "飞牌、标签、流光", "待听感复验"],
        ["三带二", "无专用报型", "飞牌、标签、流光", "音频待补"],
        ["三连对", "无专用报型", "飞牌、标签、流光", "音频待补"],
        ["钢板", "无专用报型", "飞牌、标签、流光", "音频待补"],
        ["同花顺", "授权音频已接入", "暗场、海蓝流光、中震屏", "待实机复验"],
        ["炸弹", "授权爆破声，无人声报型", "按牌数分三级表现", "待实机复验"],
        ["天王炸", "临时音效，回退炸弹声", "暗场、金色标题、强震屏", "待制作"],
        ["不要", "授权“过”语音", "纯文字淡入淡出", "已实测"],
    ]
    add_table(doc, ["牌型", "当前音频", "当前动效", "状态"], coverage, [2.35, 5.15, 5.65, 3.1], font_size=8.2)
    add_callout(doc, "命名说明", "规则要求四张王同时组成最高牌型，文档统一称为“天王炸／四王炸”，不使用斗地主的“火箭”称呼。", fill=PALE, accent=TEAL)

    # Flow + flow matrix
    doc.add_page_break()
    add_heading(doc, "2  牌局流程与反馈规格", 1)
    add_flow_strip(doc, ["大厅", "摸牌定庄", "发牌", "选牌 / 提示", "出牌循环", "结算"])
    add_heading(doc, "单机主流程", 2)
    p = doc.add_paragraph("启动 → 大厅 → 标准对局／挑战 → 摸牌定庄 → 庄家结果 → 发牌 → 循环出牌 → 本局结算 → 下一局 → 进贡／还贡。")
    style_paragraph(p, size=9.7)
    add_heading(doc, "联机分支", 2)
    p = doc.add_paragraph("大厅 → 在线对战 → 快速匹配／比赛场／好友房 → 四人齐 → 联机牌桌 → 离线／恢复。当前 Web 构建没有配置联机端点，因此只能自然复现单机牌局。")
    style_paragraph(p, size=9.7)
    flow_rows = [
        ["匹配", "无", "三张牌循环洗牌", "补进入、成功、失败短音效"],
        ["摸牌定庄", "无", "说明淡入、庄家结果", "补摸牌、翻牌、确认"],
        ["发牌", "授权发牌声", "手牌淡入、缩放、错峰", "复验联机首局"],
        ["选牌", "无", "上移 32px、金边、按压", "可加轻点击并限频"],
        ["普通出牌", "通用声＋有限报读", "克隆牌弧线飞行", "补落地层次"],
        ["倒计时", "暂用落牌声", "20 秒，末 5 秒变红", "换专用滴答声"],
        ["玩家出完", "无", "名次提示 1.6 秒", "补轻量完成音"],
        ["胜／负", "临时胜利／授权失败", "金色胜利／灰色结束", "真机校准响度"],
        ["进贡还贡", "无", "结算遮罩＋选牌", "补贡牌飞行和阶段反馈"],
        ["离线恢复", "无", "离线持续；已恢复 1.5 秒", "保持无声，不重放历史"],
    ]
    add_table(doc, ["节点", "当前声音", "当前画面", "建议"], flow_rows, [2.4, 4.1, 5.15, 4.6], font_size=8.05)

    # Screenshots 1
    doc.add_page_break()
    add_heading(doc, "3  截图材料：大厅与定庄", 1)
    add_picture_pair(doc, [
        ("01-lobby.png", "图 1  大厅：左侧局外入口，右侧主要玩法入口"),
        ("02-grouping.png", "图 2  摸牌定庄说明：目前无专用声音"),
    ])
    add_picture(doc, "03-dealer.png", "图 3  庄家结果：后续可加入一次明确的定庄确认反馈", width_inches=6.25)
    p = doc.add_paragraph("画面验收重点：局外按钮层级清楚；进入对局前的说明和结果不会与主牌桌信息混淆；定庄动效总时长建议控制在 1 秒内。")
    style_paragraph(p, size=9.3, color=MUTED)

    # Screenshots 2
    doc.add_page_break()
    add_heading(doc, "4  截图材料：牌桌与选牌", 1)
    add_picture(doc, "04-player-turn.png", "图 4  完整牌桌：左上为返回大厅和双方级数，上方队友信息独立放置", width_inches=6.35)
    add_picture_pair(doc, [
        ("05-selected-pair.png", "图 5  点按选牌：上移、金边，动作按钮动态居中"),
        ("08-hint-selection.png", "图 6  提示选牌：合法组合自动选中并显示重置／出牌"),
    ])
    p = doc.add_paragraph("选牌是最基础的交互反馈：AI 或其他玩家出牌后仍必须保持可用，取消选择应准确复位，不得因手牌重排或状态同步永久失效。")
    style_paragraph(p, size=9.3, color=MUTED)

    # Screenshots 3
    doc.add_page_break()
    add_heading(doc, "5  截图材料：跟牌与不要", 1)
    add_picture_pair(doc, [
        ("07-player-turn-follow.png", "图 7  跟牌回合：按钮和倒计时自然提示轮到自己"),
        ("06-pass-feedback.png", "图 8  不要反馈：只显示文字，不套用纸牌框"),
    ])
    add_heading(doc, "需要录屏抽帧补采的画面", 2)
    add_bullet(doc, "发牌入口动画：点击进入牌局后 100–250ms。")
    add_bullet(doc, "普通出牌飞行：点按出牌后 80–180ms。")
    add_bullet(doc, "顺子／三连对／钢板标签：约 300–380ms。")
    add_bullet(doc, "炸弹、同花顺、天王炸：分别约 520–1050ms，随机牌局难稳定复现。")
    add_bullet(doc, "玩家出完、结算和进贡还贡：需完整打完牌局。")
    add_callout(doc, "采集建议", "增加仅在开发构建启用的“牌局效果预览器”，使用固定 fixture 逐项触发。它既能生成一致截图，也能成为音效与动效回归入口。", fill=PALE, accent=TEAL)

    # Issue
    doc.add_page_break()
    add_heading(doc, "6  高优先级缺陷", 1)
    add_callout(
        doc,
        "P0 · 联机动作被误判为恢复快照",
        "GameScene.applyNetworkState() 每次收到 gameState 都先把特效计数重置为当前 playArea 长度；随后 EffectController.syncActions() 发现计数相等，直接按恢复状态返回。正常联机动作因此跳过音效、报读、飞牌、标签和主特效。",
        fill=PALE_RED,
        accent="C84E3A",
    )
    add_bullet(doc, "桌面牌仍会显示，因为 PlayAreaController 独立按最终状态渲染。")
    add_bullet(doc, "只应在首次进入、真正重连、动作序列缺口或回合恢复时 resetForRecovery()。")
    add_bullet(doc, "正常增量包应保留前一动作计数，由 syncActions() 判断并只播放最新一次动作。")
    add_heading(doc, "其他待确认", 2)
    add_bullet(doc, "顺子、同花顺、炸弹和天王炸：源码有路由，但当前构建未通过稳定听感验收。")
    add_bullet(doc, "天王炸、逢人配和胜利仍使用临时本地音效。")
    add_bullet(doc, "倒计时末 5 秒暂用落牌声，声音语义不准确。")
    add_bullet(doc, "当前没有实际背景音乐文件，音乐开关没有默认曲目。")
    add_bullet(doc, "本地“不要”有两条音效派发路径，依赖冷却避免重复，后续应统一入口。")
    add_heading(doc, "动效调度原则", 2)
    principles = [
        ["L0", "普通出牌", "飞牌、落地光环、轻触感；允许并行"],
        ["L1", "组合牌型", "短标签和扫光；不使用大面积暗场"],
        ["L2", "普通炸弹", "同一时间最多一个主特效"],
        ["L3", "同花顺／天王炸／大炸", "稀缺使用暗场、标题和震屏"],
    ]
    add_table(doc, ["强度", "适用", "调度"], principles, [2.1, 5.0, 9.2], font_size=8.5)

    # QA
    doc.add_page_break()
    add_heading(doc, "7  人工验收清单", 1)
    add_heading(doc, "音频", 2)
    for item in [
        "用手机扬声器、耳机分别测试，确保报读清晰、炸弹不过载。",
        "逐一验证已映射的单张和对子；未映射点数不能误报。",
        "顺子、同花顺、炸弹、天王炸各触发至少 3 次，记录是否漏播、重复或被截断。",
        "连续快速出牌时，通用落牌声和报读语音不能互相吞掉。",
        "“不要”只播放一次；最后 5 秒倒计时节奏均匀。",
        "关闭音效后完全静音；只关闭视觉时声音仍正常。",
    ]:
        add_bullet(doc, "□ " + item)
    add_heading(doc, "动效", 2)
    for item in [
        "四个方位都从正确起点飞向正确桌面位置。",
        "选牌上移、取消复位；其他玩家出牌后仍可继续响应。",
        "无牌可压时只显示居中的“不要”。",
        "牌型标签不遮挡桌面牌；L2/L3 不长时间遮挡必要按钮。",
        "震屏只作用于牌桌，不晃动顶部导航和弹窗。",
        "断线恢复直接显示最终状态，不重放历史飞牌和大特效。",
    ]:
        add_bullet(doc, "□ " + item)
    add_heading(doc, "联机", 2)
    for item in [
        "四个独立会话进入同房，轮流出单张、对子、不要和炸弹。",
        "其他三端都能看到飞牌、听到声音，且每个动作只触发一次。",
        "断网后其余客户端显示“离线”；恢复后 1.5 秒内显示“已恢复”。",
        "重连不会补播断线期间的全部历史声音与特效。",
    ]:
        add_bullet(doc, "□ " + item)
    add_callout(doc, "建议执行顺序", "先修复联机动作计数 → 建效果预览器 → 补缺失牌型语音 → 替换临时音效 → 真机做声音、弱网和低性能验收。", fill=PALE_GOLD, accent=GOLD)

    # Source index
    add_heading(doc, "8  源码索引", 1)
    sources = [
        ("牌面报读", "assets/scripts/audio/PlayVoiceProfiles.ts"),
        ("语义音频配置", "assets/scripts/audio/AudioProfiles.ts"),
        ("统一效果调度", "assets/scripts/effects/EffectController.ts"),
        ("牌型表现分级", "assets/scripts/effects/EffectProfileResolver.ts"),
        ("出牌飞行", "assets/scripts/effects/CardFlightController.ts"),
        ("选牌动画", "assets/scripts/ui/CardView.ts"),
        ("桌面牌与不要", "assets/scripts/ui/PlayAreaController.ts"),
        ("牌桌流程", "assets/scripts/scenes/GameScene.ts"),
        ("匹配与定庄", "assets/scripts/scenes/FrontPageController.ts"),
    ]
    add_table(doc, ["职责", "文件"], [[a, b] for a, b in sources], [4.5, 11.8], font_size=8.5)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
